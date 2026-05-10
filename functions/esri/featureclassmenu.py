# create_schema_esri.py
from kivymd.uix.screen import MDScreen
from kivymd.uix.menu import MDDropdownMenu
from kivymd.uix.button import MDButton, MDButtonText
import openpyxl
import psycopg2
import os
from kivy.app import App

from kivy.uix.widget import Widget
from kivymd.uix.dialog import MDDialog, MDDialogHeadlineText, MDDialogSupportingText, MDDialogButtonContainer
from functions.esri.featureclass_wizard import FeatureclassWizardScreen

from kivy.metrics import dp
from kivymd.uix.dialog import (
    MDDialog,
    MDDialogHeadlineText,
    MDDialogSupportingText,
    MDDialogButtonContainer,
)

class FeatureclassMenuScreen(MDScreen):

    def create_feature_dataset(self):

        # 1) Uygulamanın root’undan screen_manager’a eriş
        app = App.get_running_app()
        sm = app.root.ids.screen_manager

        # 2) Eğer sihirbaz ekranı ekli değilse ekle
        if not sm.has_screen("featureclasswizard"):
            # Wizard’ı oluşturup, seçilmiş conn+şema’yı set et
            wiz = FeatureclassWizardScreen(name="featureclasswizard")
            wiz.sel_connection = self.ids.conn_dropdown.text
            wiz.sel_schema     = self.ids.schema_dropdown.text
            sm.add_widget(wiz)


        # Sihirbaz ekranına geç
        sm.current = "featureclasswizard"

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.conn_menu = None
        self.schema_menu = None
        self.conn_info_map = {}  # Bağlantı bilgilerini burada saklıyoruz
        self.selected_connection = None  # Seçilen bağlantıyı saklamak için

    def on_pre_enter(self):
        self.load_connection_names()

    def load_connection_names(self):
        path = "documents/app_informations.xlsx"
        self.conn_names = []
        self.conn_info_map = {}

        if os.path.exists(path):
            wb = openpyxl.load_workbook(path)
            if "vt_connections" in wb.sheetnames:
                ws = wb["vt_connections"]
                headers = [cell.value for cell in ws[1]]

                for row in ws.iter_rows(min_row=2, values_only=True):
                    row_data = dict(zip(headers, row))
                    conn_name = row_data.get("Conn_Name")
                    db_type = row_data.get("DbType")

                    # Yalnızca DbType "Esri Geodatabase" olanları ekle
                    if conn_name and db_type == "Esri Geodatabase":
                        self.conn_names.append(conn_name)
                        self.conn_info_map[conn_name] = {
                            "database": row_data.get("Database"),
                            "host": row_data.get("Host"),
                            "port": row_data.get("Port"),
                            "user": row_data.get("UserName"),
                            "password": row_data.get("Password"),
                        }

        menu_items = [
            {"text": name, "on_release": lambda x=name: self.set_connection_name(x)}
            for name in self.conn_names
        ]
        self.conn_menu = MDDropdownMenu(
            caller=self.ids.conn_dropdown,
            items=menu_items,
            width_mult=3,
            position="bottom",
        )

    def open_conn_menu(self):
        if hasattr(self, 'conn_menu'):
            self.conn_menu.open()

    def set_connection_name(self, name):
        self.ids.conn_dropdown.text = name
        self.selected_connection = name  # Seçilen bağlantıyı sakla
        self.conn_menu.dismiss()
        self.build_schema_menu()  # Şemaları yükle

    def build_schema_menu(self):
        # Eğer bağlantı seçilmediyse, şema menüsünü boş bırak
        if not self.selected_connection:
            return

        info = self.conn_info_map[self.selected_connection]
        try:
            # Veritabanı bağlantısını kur
            conn = psycopg2.connect(
                database=info['database'],
                user=info['user'],
                password=info['password'],
                host=info['host'],
                port=info['port']
            )
        except Exception as e:
            self.dialog("Hata", f"Seçilen Bağlantı Bilgileri İle VT Bağlantı Sağlanaması: {str(e)}")
            return

        cursor = conn.cursor()
        cursor.execute(
            "SELECT schema_name FROM information_schema.schemata ORDER BY schema_name;"
        )
        schemas = [r[0] for r in cursor.fetchall()]
        cursor.close()

        items = [
            {
                "text": schema,
                "height": dp(48),
                "on_release": lambda x=schema: self.select_schema(x),
            }
            for schema in schemas
        ]
        self.schema_menu = MDDropdownMenu(
            caller=self.ids.schema_dropdown,
            items=items,
            width_mult=4,
        )
        self.schema_menu.open()  # Menü açılır

    def select_schema(self, schema):
        self.ids.schema_dropdown.text = schema
        self.schema_menu.dismiss()
        # Burada şemayı seçtikten sonra yapılacak işlemleri ekleyebilirsiniz
        self.ids.createfeature_btn.disabled = False  # Butonu aktif et


    def dialog(self, title, text):
        dialog= MDDialog(

            # -----------------------Headline text-------------------------
            MDDialogHeadlineText(
                text=title,
            ),
            # -----------------------Supporting text-----------------------
            MDDialogSupportingText(
                text= text,
            ),
            # ---------------------Button container------------------------
            MDDialogButtonContainer(
                Widget(),
                MDButton(
                    MDButtonText(text="Kapat"),
                    style="text",
                    on_release=lambda *args: dialog.dismiss(),  # Kapat'a basınca dialog kapanır
                ),
                spacing="8dp",
            ),
            # -------------------------------------------------------------
            # -------------------------------------------------------------
        )
        dialog.open()
        
    def dismiss_dialog(self, instance):
        if self.dialog_instance:
            self.dialog_instance.dismiss()
            self.dialog_instance = None

