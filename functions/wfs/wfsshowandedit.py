import os
import openpyxl
from kivy.factory import Factory
from kivymd.uix.screen import MDScreen
from kivymd.uix.dialog import MDDialog
from kivymd.uix.list import MDListItem, MDListItemHeadlineText
from kivymd.uix.button import MDButton, MDButtonText
from kivy.uix.widget import Widget
from kivymd.uix.dialog import (
    MDDialog,
    MDDialogHeadlineText,
    MDDialogSupportingText,
    MDDialogButtonContainer,
)

import os
import openpyxl
from kivy.metrics import dp
import requests


class WFSShowAndEditScreen(MDScreen):
    """Excel’deki bağlantıları listeler, detay gösterir ve güncellemeye izin verir."""
    data_file      = os.path.join("documents", "app_informations.xlsx")
    sheet_name     = "wfs_connections"

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.edit_mode         = False
        self.current_row_index = None
        self._dialog           = None

    def on_pre_enter(self):
        self.populate_list()
        self.clear_fields()

    def populate_list(self):
        lst = self.ids.connection_list
        lst.clear_widgets()
        if not os.path.exists(self.data_file):
            return
        wb = openpyxl.load_workbook(self.data_file)
        if self.sheet_name not in wb.sheetnames:
            return
        ws = wb[self.sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=False):
            row_idx  = row[0].row
            name     = row[1].value or ""
            text     = name
            item     = MDListItem()
            item.add_widget(MDListItemHeadlineText(text=text))
            item.bind(on_release=lambda inst, i=row_idx: self.select_connection(i))
            lst.add_widget(item)

    def select_connection(self, row_index):
        if self.edit_mode:
            return self.dialog("Hata", "Güncellemeleri Önce Kaydedin")
        wb = openpyxl.load_workbook(self.data_file)
        ws = wb[self.sheet_name]
        self.current_row_index = row_index

        # Detay alanlarını doldur
        self.ids.conn_name.text     = ws.cell(row=row_index, column=2).value or ""
        self.ids.urladress.text     = ws.cell(row=row_index, column=3).value or ""
        self.ids.username.text          = ws.cell(row=row_index, column=4).value or ""
        self.ids.password.text      = ws.cell(row=row_index, column=5).value or ""

        self.ids.wfs_update_btn.disabled = False
        self.ids.wfs_check_btn.disabled = False


    def wfs_check_connection(self):
        urladress = self.ids.urladress.text
        username= self.ids.username.text
        password = self.ids.password.text
        
        
        try:
            params = {"service": "WFS", "request": "GetCapabilities"}
            kwargs = {"params": params, "timeout": 10}

            # >>> ÖNEMLİ DEĞİŞİKLİK: username ve password boşsa auth kullanma
            if username and password:
                kwargs["auth"] = (username, password)

            response = requests.get(urladress, **kwargs)

            if response.status_code == 200 and "WFS_Capabilities" in response.text:
                self.dialog("Başarılı", "WFS bağlantısı başarılı!")

            else:
                self.dialog("Hata", f"WFS bağlantısı başarısız.\nDurum: {response.status_code}")


        except Exception as e:
            self.dialog("Hata", f"Bağlantı sağlanamadı.\nHata: {str(e)}")


    def clear_fields(self):
        for f in ("conn_name","urladress","username","password"):
            self.ids[f].text     = ""
            self.ids[f].disabled = True
        self.ids.wfs_update_btn.disabled = True
        self.ids.wfs_save_btn.disabled   = True
        self.edit_mode               = False
        self.current_row_index       = None

    def enable_editing(self):
        if not self.current_row_index:
            return
        self.edit_mode = True
        for f in ("conn_name","urladress","username","password"):
            self.ids[f].disabled = False
        self.ids.wfs_update_btn.disabled = True
        self.ids.wfs_save_btn.disabled   = False

    def save_changes(self):
        if not self.current_row_index:
            return
        wb = openpyxl.load_workbook(self.data_file)
        ws = wb[self.sheet_name]
        vals = {
            2: self.ids.conn_name.text,
            3: self.ids.urladress.text,
            4: self.ids.username.text,
            5: self.ids.password.text,
        }
        for col, val in vals.items():
            if col == 5:
                try:
                    ws.cell(row=self.current_row_index, column=col).value = int(val)
                except ValueError:
                    ws.cell(row=self.current_row_index, column=col).value = val
            else:
                ws.cell(row=self.current_row_index, column=col).value = val
        wb.save(self.data_file)

        # Düzenleme modunu kapat
        self.edit_mode = False
        for f in ("conn_name","urladress","username","password",):
            self.ids[f].disabled = True
        self.ids.wfs_update_btn.disabled = False
        self.ids.wfs_save_btn.disabled   = True


    def dialog(self, title, text):
        dialog= MDDialog(

            # -----------------------Headline text-------------------------
            MDDialogHeadlineText(
                text= "Uyarı",
            ),
            # -----------------------Supporting text-----------------------
            MDDialogSupportingText(
                text= text,
            ),
            # ---------------------Button container------------------------
            MDDialogButtonContainer(
                Widget(),
                MDButton(
                    MDButtonText(text="Tamam"),
                    style="text",
                    on_release=lambda *args: dialog.dismiss(),  # Kapat'a basınca dialog kapanır
                ),
                spacing="8dp",
            ),
            # -------------------------------------------------------------
            # -------------------------------------------------------------
        )
        dialog.open()
        
    def dismiss_dialog(self, instance):
        if self.dialog_instance:
            self.dialog_instance.dismiss()
            self.dialog_instance = None
