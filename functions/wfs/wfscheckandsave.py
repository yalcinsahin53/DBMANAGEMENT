from kivymd.uix.screen import MDScreen
from kivymd.uix.dialog import MDDialog
from kivymd.uix.button import MDButton, MDButtonText
from kivy.uix.widget import Widget
import os
import openpyxl
from kivymd.uix.dialog import (
    MDDialog,
    MDDialogHeadlineText,
    MDDialogSupportingText,
    MDDialogButtonContainer,
)
import requests


class WFSCheckScreen(MDScreen):

    def wfs_check_connection(self):
        urladress = self.ids.urladress.text
        username= self.ids.username.text
        password = self.ids.password.text
        
        try:
            params = {"service": "WFS", "request": "GetCapabilities"}
            kwargs = {"params": params, "timeout": 10}

            # >>> ÖNEMLİ DEĞİŞİKLİK: username ve password boşsa auth kullanma
            if username and password:
                kwargs["auth"] = (username, password)

            response = requests.get(urladress, **kwargs)

            if response.status_code == 200 and "WFS_Capabilities" in response.text:
                self.dialog("Başarılı", "WFS bağlantısı başarılı!")
                self.ids.wfs_save_button.disabled = False
            else:
                self.dialog("Hata", f"WFS bağlantısı başarısız.\nDurum: {response.status_code}")
                self.ids.wfs_save_button.disabled = True

        except Exception as e:
            self.dialog("Hata", f"Bağlantı sağlanamadı.\nHata: {str(e)}")
            self.ids.wfs_save_button.disabled = True

    
    def wfs_save_connection(self):
        conn_name = self.ids.conn_name.text
        urladress = self.ids.urladress.text
        username = self.ids.username.text
        password = self.ids.password.text
        
        try:
            if not os.path.exists("documents"):
                os.makedirs("documents")
            file_path = os.path.join("documents", "app_informations.xlsx")

            if not os.path.exists(file_path):
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "wfs_connections"
                ws.append(["No", "Conn_Name", "UrlAdress", "UserName", "Password",])
            else:
                wb = openpyxl.load_workbook(file_path)
                if "wfs_connections" in wb.sheetnames:
                    ws = wb["wfs_connections"]
                else:
                    ws = wb.create_sheet("wfs_connections")
                    ws.append(["No", "Conn_Name",  "UrlAdress", "UserName", "Password",])

            next_no = ws.max_row if ws.max_row > 1 else 1
            ws.append([next_no, conn_name, urladress, username, password, ])
            wb.save(file_path)

            self.dialog("Başarılı", "Bilgiler kaydedildi!")
            self.ids.wfs_save_button.disabled = True
        except Exception as e:
            self.dialog("Hata", f"Kayıt yapılamadı.\nHata: {str(e)}")

    def dialog(self, title, text):
        dialog= MDDialog(

            # -----------------------Headline text-------------------------
            MDDialogHeadlineText(
                text=title,
            ),
            # -----------------------Supporting text-----------------------
            MDDialogSupportingText(
                text= text,
            ),
            # ---------------------Button container------------------------
            MDDialogButtonContainer(
                Widget(),
                MDButton(
                    MDButtonText(text="Kapat"),
                    style="text",
                    on_release=lambda *args: dialog.dismiss(),  # Kapat'a basınca dialog kapanır
                ),
                spacing="8dp",
            ),
            # -------------------------------------------------------------
            # -------------------------------------------------------------
        )
        dialog.open()
        
    def dismiss_dialog(self):
        if self.dialog_instance:
            self.dialog_instance.dismiss()
            self.dialog_instance = None
